#!/usr/bin/env python3
"""把一个目录下的图片按文件名顺序组装成 16:9 全幅图 PPTX。

用法: python3 build_pptx.py <图片目录> <输出.pptx>
图片命名建议: 01-封面.png, 02-xxx.png ...（按文件名排序）
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

CANVAS_W = Emu(12192000)  # 16:9
CANVAS_H = Emu(6858000)
EXTS = {".png", ".jpg", ".jpeg", ".webp"}


def main() -> None:
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    img_dir, out_path = Path(sys.argv[1]), Path(sys.argv[2])
    images = sorted(p for p in img_dir.iterdir() if p.suffix.lower() in EXTS)
    if not images:
        sys.exit(f"目录里没有图片: {img_dir}")

    prs = Presentation()
    prs.slide_width, prs.slide_height = CANVAS_W, CANVAS_H
    blank = prs.slide_layouts[6]
    for img in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=CANVAS_W, height=CANVAS_H)
    prs.save(out_path)
    print(f"OK: {len(images)} 页 -> {out_path}")


if __name__ == "__main__":
    main()
